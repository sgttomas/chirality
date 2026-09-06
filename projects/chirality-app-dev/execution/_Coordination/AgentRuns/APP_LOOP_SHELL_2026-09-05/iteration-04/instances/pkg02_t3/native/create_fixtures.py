from pathlib import Path
import json, hashlib, sys
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen import canvas
root = Path('/private/tmp/app-loop-t3-native-ac3dexfs/fixtures')
root.mkdir(exist_ok=True)
(root / 'Readable note.md').write_text('# T3 fixture document\n\nActual local file endpoint proof.\n\n- First item\n- Second item\n\n```ts\nconst fileView = true;\n```\n')
(root / 'plain.txt').write_text('T3 plain text fixture: no authority or runtime content.\n')
(root / 'data.csv').write_text('name,value\nT3,42\n')
(root / 'oversized.txt').write_bytes(b'x' * (10 * 1024 * 1024 + 1))
(root / 'unsupported.bin').write_bytes(bytes(range(32)))
pdf = canvas.Canvas(str(root / 'Preview.pdf'))
pdf.drawString(60, 760, 'T3 native PDF fixture')
pdf.drawString(60, 730, 'Actual document display must show this text.')
pdf.save()
doc = Document(); doc.add_heading('T3 Quick Look DOCX fixture', 0); doc.add_paragraph('Native document preview evidence.'); doc.save(root / 'Preview.docx')
book = Workbook(); sheet = book.active; sheet['A1'] = 'T3 Quick Look XLSX fixture'; sheet['A2'] = 'Value'; sheet['B2'] = 42; sheet.column_dimensions['A'].width = 40; book.save(root / 'Preview.xlsx')
deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[5]); slide.shapes.title.text = 'T3 Quick Look PPTX fixture'; slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(1)).text_frame.text = 'Native slide preview evidence.'; deck.save(root / 'Preview.pptx')
longdir = root / ('Long path ' + 'document provenance ' * 5); longdir.mkdir(exist_ok=True); (longdir / 'Long filename requiring bounded wrapping.md').write_text('# Bounded long path\n\nThe panel must remain readable.\n')
files = [{"path":str(p),"bytes":p.stat().st_size,"sha256":hashlib.sha256(p.read_bytes()).hexdigest()} for p in sorted(root.rglob('*')) if p.is_file()]
print(json.dumps({"python":sys.version,"root":str(root),"files":files},indent=2))
